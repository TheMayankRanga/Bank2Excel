from flask import Flask, request, render_template, jsonify, send_file
from pathlib import Path
from datetime import datetime
import io, re, base64, shutil, os

import fitz
import pytesseract
from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

APP_NAME = "Bank2Excel"
APP_TAGLINE = "Bank Statement Converter"

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024

# Works locally on Windows and automatically uses the Linux Tesseract binary
# installed by the Docker image when deployed.
TESSERACT_CANDIDATES = [
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
    r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    os.path.expandvars(r"%LOCALAPPDATA%\Programs\Tesseract-OCR\tesseract.exe"),
    "/usr/bin/tesseract",
]
for _t in TESSERACT_CANDIDATES:
    if os.path.exists(_t):
        pytesseract.pytesseract.tesseract_cmd = _t
        break

DATE_RE = re.compile(
    r"^\s*("
    r"\d{1,2}[/-]\d{1,2}(?:[/-]\d{2,4})?"
    r"|\d{1,2}\s+[A-Za-z]{3,9}\s+\d{2,4}"
    r"|[A-Za-z]{3,9}\s+\d{1,2}(?:,)?\s+\d{2,4}"
    r")\b"
)
AMOUNT_RE = re.compile(
    r"(?<![\w/])(?:\(?-?\$?\s?(?:(?:\d{1,3}(?:,\d{3})*(?:\.\d{1,2})?)|(?:\d+(?:\.\d{1,2})?)|(?:\.\d{1,2})))(?:\)?)(?!\w)"
)
IGNORE_RE = re.compile(
    r"^(page\s+\d+|account\s*(number|#)?\b|statement\b|summary\b|"
    r"beginning balance\b|ending balance\b)", re.I
)
SECTION_RESET_RE = re.compile(
    r"^(deposits? and other credits?|withdrawals? and other debits?|"
    r"service charges? and fees?|checks? paid|electronic withdrawals?|"
    r"electronic deposits?|transactions? by date|account activity|activity for|"
    r"transaction history|transaction details?)\b", re.I
)
CREDIT_WORDS = re.compile(
    r"\b(credit|deposit|salary|payroll|refund|reversal|interest\s*credit|"
    r"cash\s*deposit|received|transfer\s*in|incoming|cr)\b", re.I
)
DEBIT_WORDS = re.compile(
    r"\b(debit|withdraw|withdrawal|purchase|pos|atm|fee|charge|payment|"
    r"check|cheque|transfer\s*out|outgoing|dr)\b", re.I
)


def clean_amount(s):
    s = s.strip().replace("$", "").replace(",", "").replace(" ", "")
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    try:
        v = float(s)
        return -v if neg else v
    except Exception:
        return None


def extract_native_pages(pdf_bytes):
    pages = []
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    try:
        for page in doc:
            pages.append(page.get_text("text") or "")
    finally:
        doc.close()
    return pages


def extract_ocr_pages(pdf_bytes):
    pages = []
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    try:
        for page in doc:
            pix = page.get_pixmap(matrix=fitz.Matrix(2.4, 2.4), alpha=False)
            img = Image.open(io.BytesIO(pix.tobytes("png")))
            pages.append(pytesseract.image_to_string(img, config="--psm 6") or "")
    finally:
        doc.close()
    return pages


def extract_native(pdf_bytes):
    return "\n".join(extract_native_pages(pdf_bytes))


def extract_ocr(pdf_bytes):
    return "\n".join(extract_ocr_pages(pdf_bytes))


def looks_like_statement(text):
    low = text.lower()
    return any(k in low for k in ("date", "description", "debit", "credit", "withdrawal", "deposit", "balance", "transaction"))


def normalize_lines(text):
    return [re.sub(r"\s+", " ", x).strip() for x in text.splitlines() if x.strip()]


def extract_statement_years(text):
    years = []
    for m in re.finditer(r"(?<!\d)(20\d{2}|19\d{2})(?!\d)", text):
        y = int(m.group(1))
        if 1990 <= y <= datetime.now().year + 1:
            years.append(y)
    return sorted(set(years))


def find_statement_year(text):
    """Prefer years near statement-period words; otherwise use the most useful year found."""
    years = extract_statement_years(text)
    if not years:
        return None

    lines = normalize_lines(text)
    for i, line in enumerate(lines):
        low = line.lower()
        if any(k in low for k in ("statement period", "statement date", "period ending", "through", "from", "to")):
            nearby = " ".join(lines[max(0, i - 1): i + 2])
            m = re.search(r"(19\d{2}|20\d{2})", nearby)
            if m:
                return int(m.group(1))
    return years[0]



def extract_pdf_metadata_year(pdf_bytes):
    """Use the PDF's own creation/modification metadata as another year signal."""
    try:
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        try:
            meta = doc.metadata or {}
            for key in ("creationDate", "modDate"):
                value = meta.get(key) or ""
                m = re.search(r"(19\d{2}|20\d{2})", value)
                if m:
                    year = int(m.group(1))
                    if 1990 <= year <= datetime.now().year + 1:
                        return year
        finally:
            doc.close()
    except Exception:
        pass
    return None


def infer_year(pdf_bytes, filename, text, year_hint):
    """Return a useful year without forcing the user to type one."""
    if year_hint:
        return int(year_hint), "manual"
    statement_year = find_statement_year(text)
    if statement_year:
        return statement_year, "PDF"
    m = re.search(r"(?<!\d)(19\d{2}|20\d{2})(?!\d)", filename or "")
    if m:
        return int(m.group(1)), "filename"
    metadata_year = extract_pdf_metadata_year(pdf_bytes)
    if metadata_year:
        return metadata_year, "PDF metadata"
    return datetime.now().year, "automatic fallback"

def parse_date_parts(date_text):
    s = date_text.strip().replace(",", "")
    m = re.match(r"^(\d{1,2})[/-](\d{1,2})(?:[/-](\d{2,4}))?$", s)
    if m:
        month, day = int(m.group(1)), int(m.group(2))
        year = int(m.group(3)) if m.group(3) else None
        if year is not None and year < 100:
            year += 2000
        if 1 <= month <= 12 and 1 <= day <= 31:
            return month, day, year
        return None

    m = re.match(r"^(\d{1,2})\s+([A-Za-z]{3,9})\s+(\d{2,4})$", s, re.I)
    if m:
        try:
            month = datetime.strptime(m.group(2)[:3].title(), "%b").month
        except ValueError:
            return None
        year = int(m.group(3))
        return month, int(m.group(1)), year + 2000 if year < 100 else year

    m = re.match(r"^([A-Za-z]{3,9})\s+(\d{1,2})\s+(\d{2,4})$", s, re.I)
    if m:
        try:
            month = datetime.strptime(m.group(1)[:3].title(), "%b").month
        except ValueError:
            return None
        year = int(m.group(3))
        return month, int(m.group(2)), year + 2000 if year < 100 else year
    return None


def format_full_date(date_text, statement_year=None, previous_full_date=None, year_hint=None):
    parts = parse_date_parts(date_text)
    if not parts:
        return date_text

    month, day, explicit_year = parts
    if explicit_year is not None:
        return f"{month:02d}/{day:02d}/{explicit_year:04d}"

    base_year = int(year_hint or statement_year or datetime.now().year)
    prev = None
    if previous_full_date:
        try:
            prev = datetime.strptime(previous_full_date, "%m/%d/%Y")
        except ValueError:
            pass

    year = prev.year if prev else base_year
    # Detect a statement crossing New Year: Dec 30 -> Jan 02.
    if prev:
        if month < prev.month and (prev.month - month) >= 6:
            year += 1
        elif month > prev.month and (month - prev.month) >= 10:
            year -= 1
    return f"{month:02d}/{day:02d}/{year:04d}"


def focus_primary_transaction_tables(pages):
    focused = []
    for page in pages:
        lines = normalize_lines(page)
        top = " ".join(lines[:18]).lower()
        if "by type" in top or "detailed description" in top:
            continue

        header_idx = None
        for i, line in enumerate(lines):
            low = line.lower()
            if "date" in low and "description" in low and any(k in low for k in ("debit", "credit", "balance", "amount", "withdrawal", "deposit")):
                header_idx = i
                break

        if header_idx is not None:
            block = [lines[header_idx]]
            for line in lines[header_idx + 1:]:
                low = line.lower()
                if "account transactions by type" in low or "detailed description" in low or SECTION_RESET_RE.match(line):
                    break
                block.append(line)
            focused.append("\n".join(block))
            continue

        if "continued" in top and "transaction" in top:
            block, started = [], False
            for line in lines:
                if DATE_RE.match(line):
                    started = True
                if started:
                    if SECTION_RESET_RE.match(line):
                        break
                    block.append(line)
            if block:
                focused.append("\n".join(block))
    return "\n".join(focused) if focused else ""


def choose_transaction_amount(tx, balance, previous_balance, desc):
    if not tx:
        return None
    if balance is not None and previous_balance is not None:
        movement = round(abs(balance - previous_balance), 2)
        matches = [x for x in tx if round(abs(x), 2) == movement]
        if len(matches) == 1:
            return matches[0]
        scaled = []
        for x in tx:
            for divisor in (10, 100, 1000):
                if round(abs(x) / divisor, 2) == movement:
                    scaled.append(x / divisor)
        if len(scaled) == 1:
            return scaled[0]
    if len(tx) >= 2 and re.search(r"\b(check|cheque)\b", desc, re.I):
        return tx[-1]
    return tx[-1] if len(tx) == 1 else None


def classify_amount(chosen, balance, previous_balance, desc):
    debit = credit = None
    dr = bool(re.search(r"\bdr\b", desc, re.I))
    cr = bool(re.search(r"\bcr\b", desc, re.I))
    if dr and not cr:
        debit = abs(chosen)
    elif cr and not dr:
        credit = abs(chosen)
    elif balance is not None and previous_balance is not None:
        if balance > previous_balance:
            credit = abs(chosen)
        elif balance < previous_balance:
            debit = abs(chosen)
    elif CREDIT_WORDS.search(desc) and not DEBIT_WORDS.search(desc):
        credit = abs(chosen)
    else:
        debit = abs(chosen)
    return debit, credit


def parse_rows(text, year_hint=None):
    lines = normalize_lines(text)
    statement_year = find_statement_year(text)
    rows, seen = [], set()
    previous_balance = None
    previous_date = None
    header_mode = ""

    for line in lines:
        low = line.lower()
        if SECTION_RESET_RE.match(line):
            previous_balance = None
            continue
        if "date" in low and any(k in low for k in ("description", "debit", "credit", "balance", "withdrawal", "deposit", "amount")):
            header_mode = low
            continue
        if IGNORE_RE.match(line):
            continue

        m = DATE_RE.match(line)
        if not m:
            continue
        raw_date = m.group(1)
        date = format_full_date(raw_date, statement_year, previous_date, year_hint)
        rest = line[m.end():].strip()

        raw_amounts = AMOUNT_RE.findall(rest)
        amounts, raw_kept = [], []
        for raw in raw_amounts:
            value = clean_amount(raw)
            if value is not None:
                amounts.append(value)
                raw_kept.append(raw)
        if not amounts:
            continue

        balance = amounts[-1] if len(amounts) >= 2 else None
        tx = amounts[:-1] if len(amounts) >= 2 else amounts[:]
        chosen = choose_transaction_amount(tx, balance, previous_balance, rest)

        desc = rest
        remove_raw = []
        if raw_kept:
            if len(amounts) >= 2:
                remove_raw.append(raw_kept[-1])
            if chosen is not None:
                for raw, value in zip(raw_kept[:len(tx)], tx):
                    if round(abs(value), 2) == round(abs(chosen), 2) or any(round(abs(value) / d, 2) == round(abs(chosen), 2) for d in (10, 100, 1000)):
                        remove_raw.append(raw)
                        break
            elif len(tx) >= 2:
                remove_raw.extend(raw_kept[-len(tx):])
            elif len(tx) == 1:
                remove_raw.append(raw_kept[0])
        for raw in remove_raw:
            desc = desc.replace(raw, " ", 1)
        desc = re.sub(r"\s+", " ", desc).strip(" -:|.") or "Transaction"

        debit = credit = None
        if chosen is not None:
            debit, credit = classify_amount(chosen, balance, previous_balance, desc)
        elif len(tx) >= 2:
            if "debit" in header_mode or "withdraw" in header_mode:
                debit, credit = abs(tx[-2]), abs(tx[-1])
            elif CREDIT_WORDS.search(desc) and not DEBIT_WORDS.search(desc):
                credit, debit = abs(tx[-1]), abs(tx[-2])
            else:
                debit, credit = abs(tx[-2]), abs(tx[-1])
        elif len(tx) == 1:
            if CREDIT_WORDS.search(desc) and not DEBIT_WORDS.search(desc):
                credit = abs(tx[0])
            else:
                debit = abs(tx[0])

        if balance is not None:
            previous_balance = balance
        previous_date = date

        if desc.lower() in {"date", "description", "amount"}:
            continue

        amount = debit if debit is not None else credit
        desc_key = re.sub(r"[^a-z0-9]+", "", desc.lower())
        key = (date, round(amount or 0, 2), round(balance, 2) if balance is not None else desc_key)
        if key in seen:
            continue
        seen.add(key)
        rows.append({"date": date, "description": desc, "debit": debit, "credit": credit, "balance": balance})

    return rows


def make_xlsx(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(["Date", "Description", "Debit", "Credit", "Balance"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2563EB")
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        try:
            excel_date = datetime.strptime(r["date"], "%m/%d/%Y")
        except Exception:
            excel_date = r["date"]
        ws.append([excel_date, r["description"], r["debit"], r["credit"], r["balance"]])
    for cell in ws["A"][1:]:
        cell.number_format = "mm/dd/yyyy"
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=5):
        for cell in row:
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = "#,##0.00"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, width in enumerate([16, 60, 16, 16, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


@app.get("/")
def home():
    return render_template("index.html", app_name=APP_NAME, app_tagline=APP_TAGLINE)


@app.get("/health")
def health():
    return jsonify(status="ok")


@app.post("/api/convert")
def convert():
    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify(error="Please choose a PDF."), 400
    if not f.filename.lower().endswith(".pdf"):
        return jsonify(error="Only PDF files are supported."), 400
    data = f.read()
    if not data:
        return jsonify(error="The PDF is empty."), 400

    try:
        year_raw = (request.form.get("year") or "").strip()
        year_hint = None
        if year_raw:
            if not year_raw.isdigit() or not (1900 <= int(year_raw) <= 2100):
                return jsonify(error="Please enter a valid statement year between 1900 and 2100."), 400
            year_hint = int(year_raw)

        native_pages = extract_native_pages(data)
        text = "\n".join(native_pages)
        method = "PDF text"
        ocr_pages = None
        if len(text.strip()) < 80 or not looks_like_statement(text):
            tesseract_path = pytesseract.pytesseract.tesseract_cmd
            if not (shutil.which("tesseract") or os.path.exists(tesseract_path)):
                return jsonify(error="OCR is not installed on this server. Use the deployment package or install Tesseract locally."), 422
            ocr_pages = extract_ocr_pages(data)
            text = "\n".join(ocr_pages)
            method = "OCR"

        detected_year, year_source = infer_year(data, f.filename, text, year_hint)
        pages_for_focus = ocr_pages if ocr_pages is not None else native_pages
        focused_text = focus_primary_transaction_tables(pages_for_focus)
        parse_text = focused_text or text
        rows = parse_rows(parse_text, year_hint=detected_year)
        if not rows:
            return jsonify(error="The PDF was readable, but no transaction rows could be identified. Try another statement or a higher-quality scan.", method=method), 422

        xlsx = make_xlsx(rows)
        encoded = base64.b64encode(xlsx.getvalue()).decode("ascii")
        safe_name = Path(f.filename).stem or "bank_statement"
        return jsonify(
            rows=rows,
            count=len(rows),
            method=method,
            filename=safe_name + "_converted.xlsx",
            xlsx=encoded,
            year_used=detected_year,
            year_source=year_source,
        )
    except Exception as e:
        return jsonify(error=f"Could not process this PDF: {e}"), 500


@app.errorhandler(413)
def too_large(_):
    return jsonify(error="PDF is too large. Maximum size is 25 MB."), 413


if __name__ == "__main__":
    print(f"{APP_NAME} running at http://127.0.0.1:8000")
    app.run(host="127.0.0.1", port=8000, debug=False)
